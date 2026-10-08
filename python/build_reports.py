#!/usr/bin/env python3
"""Standard monthly MI pack (Excel, live formulas) + regional packs + an e-mail-ready distribution file."""
import sqlite3, subprocess, sys
from email.message import EmailMessage
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, str(Path(__file__).parent))
import insights as ins

ROOT = Path(__file__).resolve().parents[1]
RECALC = Path("/mnt/skills/public/xlsx/scripts/recalc.py")
F = "Arial"
HDR = PatternFill("solid", fgColor="1F3A4D"); BAND = PatternFill("solid", fgColor="EAF0F4")
thin = Side(style="thin", color="C9D3DA")


def hdr(ws, row, ncols, start=1):
    for c in range(start, start + ncols):
        x = ws.cell(row, c); x.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        x.fill = HDR; x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def put_df(ws, df, r0=1, fmts=None):
    for j, c in enumerate(df.columns, 1): ws.cell(r0, j, c)
    hdr(ws, r0, len(df.columns))
    for i, row in enumerate(df.itertuples(index=False), r0 + 1):
        for j, v in enumerate(row, 1):
            if pd.isna(v): v = None
            ws.cell(i, j, v).font = Font(name=F, size=10)
    for j, c in enumerate(df.columns, 1):
        if fmts and c in fmts:
            for i in range(r0 + 1, r0 + 1 + len(df)): ws.cell(i, j).number_format = fmts[c]
        ws.column_dimensions[L(j)].width = max(11, min(34, len(str(c)) + 4))
    ws.freeze_panes = ws.cell(r0 + 1, 1)


def fetch(con, region=None):
    w = "" if region is None else f"WHERE region_name = '{region}'"
    t = pd.read_sql(f"SELECT * FROM v_treaty_summary {w} ORDER BY treaty_id", con)
    wl = "" if region is None else f"AND region_name = '{region}'"
    large = pd.read_sql(f"SELECT * FROM v_large_losses WHERE 1=1 {wl} ORDER BY incurred_amount DESC LIMIT 25", con)
    watch = pd.read_sql(f"SELECT treaty_id, cedant_name, lob_name, region_name, underwriter_name, limit_usd, incurred, limit_utilisation "
                        f"FROM v_watchlist WHERE 1=1 {wl} ORDER BY limit_utilisation DESC", con)
    rj = "" if region is None else f"WHERE r.region_name = '{region}'"
    monthly = pd.read_sql(f"""
      WITH m AS (SELECT DISTINCT year_month ym FROM dim_date WHERE year_month BETWEEN '2021-01' AND '2026-09'),
      p AS (SELECT substr(f.txn_date,1,7) ym, SUM(f.gross_written_premium) gwp, SUM(f.earned_premium) earned FROM fact_premium f
            JOIN dim_treaty t ON t.treaty_id=f.treaty_id JOIN dim_region r ON r.region_id=t.region_id {rj} GROUP BY 1),
      c AS (SELECT substr(f.loss_date,1,7) ym, COUNT(*) claims, SUM(f.incurred_amount) incurred FROM fact_claims f
            JOIN dim_treaty t ON t.treaty_id=f.treaty_id JOIN dim_region r ON r.region_id=t.region_id {rj} GROUP BY 1)
      SELECT m.ym year_month, COALESCE(p.gwp,0) gwp, COALESCE(p.earned,0) earned, COALESCE(c.claims,0) claims,
             COALESCE(c.incurred,0) incurred FROM m LEFT JOIN p ON p.ym=m.ym LEFT JOIN c ON c.ym=m.ym ORDER BY 1""", con)
    return t, large, watch, monthly


def build_workbook(con, path, region=None):
    t, large, watch, monthly = fetch(con, region)
    dq = pd.read_sql("SELECT check_id, check_name, category, severity, records_tested, records_failed, status, action_taken FROM dq_check_log ORDER BY check_id", con)
    wb = Workbook(); S = wb.active; S.title = "Summary"
    # ---- data sheets first (Summary formulas point at them) ----
    D = wb.create_sheet("Treaty_Data")
    cols = ["treaty_id", "cedant_name", "lob_name", "region_name", "underwriter_name", "treaty_type", "underwriting_year",
            "limit_usd", "gwp", "ceded", "earned", "earned_commission", "claim_count", "paid", "incurred", "exposure", "risk_band"]
    d = t[cols].copy()
    put_df(D, d, fmts={c: "#,##0" for c in ["limit_usd", "gwp", "ceded", "earned", "earned_commission", "paid", "incurred", "exposure"]})
    D.cell(1, 18, "loss_ratio"); hdr(D, 1, 1, 18)
    n = len(d) + 1
    for i in range(2, n + 1):
        D.cell(i, 18, f"=IFERROR(O{i}/K{i},0)").number_format = "0.0%"; D.cell(i, 18).font = Font(name=F, size=10)
    D.auto_filter.ref = f"A1:R{n}"
    rng = lambda col: f"Treaty_Data!${col}$2:${col}${n}"

    Mt = wb.create_sheet("Monthly_Trend")
    put_df(Mt, monthly, fmts={"gwp": "#,##0", "earned": "#,##0", "incurred": "#,##0"})
    for c, h in [(6, "gwp_mom_growth"), (7, "incurred_rolling_3m")]: Mt.cell(1, c, h)
    hdr(Mt, 1, 2, 6)
    for i in range(2, len(monthly) + 2):
        Mt.cell(i, 6, f"=IFERROR(B{i}/B{i-1}-1,\"\")" if i > 2 else None).number_format = "0.0%"
        Mt.cell(i, 7, f"=AVERAGE(E{max(2,i-2)}:E{i})").number_format = "#,##0"
    ch = LineChart(); ch.title = "Monthly gross premium vs incurred losses (USD)"; ch.height, ch.width = 8, 22
    ch.add_data(Reference(Mt, min_col=2, min_row=1, max_row=len(monthly) + 1), titles_from_data=True)
    ch.add_data(Reference(Mt, min_col=5, min_row=1, max_row=len(monthly) + 1), titles_from_data=True)
    ch.set_categories(Reference(Mt, min_col=1, min_row=2, max_row=len(monthly) + 1))
    Mt.add_chart(ch, "I2")

    LL = wb.create_sheet("Large_Losses")
    put_df(LL, large, fmts={c: "#,##0" for c in ["paid_amount", "case_reserve", "incurred_amount"]})
    W = wb.create_sheet("Watchlist")
    put_df(W, watch, fmts={"limit_usd": "#,##0", "incurred": "#,##0", "limit_utilisation": "0%"})
    DQ = wb.create_sheet("DQ_Log"); put_df(DQ, dq)
    DQ.conditional_formatting.add(f"G2:G{len(dq)+1}", CellIsRule(operator="equal", formula=['"FAIL"'], font=Font(color="9C0006", bold=True), fill=PatternFill("solid", bgColor="FFC7CE")))
    DQ.conditional_formatting.add(f"G2:G{len(dq)+1}", CellIsRule(operator="equal", formula=['"PASS"'], font=Font(color="006100"), fill=PatternFill("solid", bgColor="C6EFCE")))
    for col, w in {"B": 58, "H": 26}.items(): DQ.column_dimensions[col].width = w
    NT = wb.create_sheet("Notes")

    # ---- Summary ----
    title = "Reinsurance Portfolio MI" + (f" - {region}" if region else " - Total Portfolio")
    S["A1"] = title; S["A1"].font = Font(name=F, size=16, bold=True, color="1F3A4D")
    S["A2"] = "Valuation date 30-Sep-2026 | USD millions unless stated | all figures are live formulas over Treaty_Data"
    S["A2"].font = Font(name=F, size=9, italic=True)
    S["A4"], S["B4"] = "Admin expense ratio (assumption)", 0.05
    S["B4"].font = Font(name=F, color="0000FF", bold=True); S["B4"].number_format = "0.0%"; S["B4"].fill = PatternFill("solid", fgColor="FFFF00")
    S["C4"] = "Input: edit to flex the combined ratio. Source: ref_assumptions table."; S["C4"].font = Font(name=F, size=9, italic=True)
    kp = [("Gross written premium", f"=SUM({rng('I')})/1000000", "#,##0.0"), ("Ceded premium", f"=SUM({rng('J')})/1000000", "#,##0.0"),
          ("Earned premium", f"=SUM({rng('K')})/1000000", "#,##0.0"), ("Incurred losses", f"=SUM({rng('O')})/1000000", "#,##0.0"),
          ("Loss ratio", "=IFERROR(D7/C7,0)", "0.0%"), ("Combined ratio", f"=IFERROR(E7+SUM({rng('L')})/1000000/C7+B4,0)", "0.0%"),
          ("Limit exposure", f"=SUM({rng('P')})/1000000", "#,##0"), ("Treaties", f"=COUNTA({rng('A')})", "#,##0"),
          ("Claims", f"=SUM({rng('M')})", "#,##0")]
    heads = ["Gross premium", "Ceded", "Earned", "Incurred", "Loss ratio", "Combined ratio", "Exposure", "Treaties", "Claims"]
    for j, (h, (_, f, nf)) in enumerate(zip(heads, kp), 1):
        S.cell(6, j, h); S.cell(7, j, f).number_format = nf
        S.cell(7, j).font = Font(name=F, size=14, bold=True, color="1F3A4D")
        S.cell(7, j).alignment = Alignment(horizontal="center")
    hdr(S, 6, 9)
    # data health
    S["K6"], S["K7"] = "Data health", f'=COUNTIF(DQ_Log!G2:G{len(dq)+1},"PASS")&" of "&COUNTA(DQ_Log!G2:G{len(dq)+1})&" checks passed"'
    hdr(S, 6, 1, 11); S["K7"].font = Font(name=F, bold=True, color="9C0006"); S.column_dimensions["K"].width = 26

    def block(r0, label, keycol, keys, fmt_key=None):
        for j, h in enumerate([label, "GWP", "Earned", "Incurred", "Loss ratio", "Commission ratio", "Combined ratio", "Treaties"], 1): S.cell(r0, j, h)
        hdr(S, r0, 8)
        for i, k in enumerate(keys, r0 + 1):
            S.cell(i, 1, k)
            crit = f'$A{i}' if not isinstance(k, (int,)) else f'$A{i}'
            S.cell(i, 2, f"=SUMIFS({rng('I')},{rng(keycol)},{crit})/1000000").number_format = "#,##0.0"
            S.cell(i, 3, f"=SUMIFS({rng('K')},{rng(keycol)},{crit})/1000000").number_format = "#,##0.0"
            S.cell(i, 4, f"=SUMIFS({rng('O')},{rng(keycol)},{crit})/1000000").number_format = "#,##0.0"
            S.cell(i, 5, f"=IFERROR(D{i}/C{i},0)").number_format = "0.0%"
            S.cell(i, 6, f"=IFERROR(SUMIFS({rng('L')},{rng(keycol)},{crit})/1000000/C{i},0)").number_format = "0.0%"
            S.cell(i, 7, f"=E{i}+F{i}+$B$4").number_format = "0.0%"
            S.cell(i, 8, f"=COUNTIF({rng(keycol)},{crit})").number_format = "#,##0"
            for j in range(1, 9): S.cell(i, j).font = Font(name=F, size=10)
        last = r0 + len(keys)
        S.conditional_formatting.add(f"G{r0+1}:G{last}", CellIsRule(operator="greaterThan", formula=["1"], font=Font(color="9C0006", bold=True), fill=PatternFill("solid", bgColor="FFC7CE")))
        S.conditional_formatting.add(f"E{r0+1}:E{last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1.2, color="7FA6BF"))
        return last
    lobs = sorted(d.lob_name.unique()); regs = sorted(d.region_name.unique()); uwys = sorted(int(x) for x in d.underwriting_year.unique())
    e1 = block(10, "Line of business", "C", lobs)
    e2 = block(e1 + 3, "Region", "D", regs)
    e3 = block(e2 + 3, "Underwriting year", "G", uwys)
    bc = BarChart(); bc.type = "col"; bc.title = "Loss ratio by line of business"; bc.height, bc.width = 7, 13; bc.legend = None
    bc.add_data(Reference(S, min_col=5, min_row=10, max_row=e1), titles_from_data=True)
    bc.set_categories(Reference(S, min_col=1, min_row=11, max_row=e1)); S.add_chart(bc, f"J{10}")
    S.column_dimensions["A"].width = 30
    for c in "BCDEFGHI": S.column_dimensions[c].width = 15
    S.sheet_view.showGridLines = False

    notes = ["Definitions", "Gross written premium (GWP): premium booked from cedants, converted to USD at valuation-date FX.",
             "Earned premium: portion of GWP for risk exposure elapsed to the valuation date (pro-rata by month).",
             "Ceded premium: premium passed on to retrocessionaires.", "Incurred = paid + case reserve (no IBNR in this dataset).",
             "Loss ratio = incurred / earned premium. Combined ratio = loss ratio + earned commission ratio + admin expense assumption.",
             "Underwriting year = year the treaty incepted. Recent years are immature: open claims still develop.",
             "Quarantined source records (see DQ_Log) are excluded. Unmapped line-of-business codes appear as 'Unmapped'.",
             "Source: SQLite warehouse built by python/pipeline.py; views in sql/04_mi_views.sql."]
    for i, x in enumerate(notes, 1):
        NT.cell(i, 1, x).font = Font(name=F, size=10, bold=(i == 1))
    NT.column_dimensions["A"].width = 130
    wb.save(path)
    if RECALC.exists():
        r = subprocess.run([sys.executable, str(RECALC), str(path), "120"], capture_output=True, text=True)
        if '"status": "success"' not in r.stdout: print("RECALC WARNING:", r.stdout[-400:], r.stderr[-200:])
    return path


def email(path, body, out):
    m = EmailMessage(); m["Subject"] = "Monthly Reinsurance Portfolio MI - 30 Sep 2026"
    m["To"] = "underwriting-team@example.com"; m["From"] = "portfolio.analytics@example.com"
    m.set_content(body)
    m.add_attachment(Path(path).read_bytes(), maintype="application",
                     subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=Path(path).name)
    Path(out).write_bytes(bytes(m))


def main(root=ROOT):
    root = Path(root); con = sqlite3.connect(root / "data" / "reinsurance.db")
    (root / "reports" / "regional").mkdir(parents=True, exist_ok=True); (root / "distribution").mkdir(exist_ok=True)
    main_pack = build_workbook(con, root / "reports" / "Monthly_MI_Pack.xlsx")
    for r in pd.read_sql("SELECT region_name FROM dim_region WHERE region_name <> 'Unassigned'", con).region_name:
        build_workbook(con, root / "reports" / "regional" / f"MI_{r.replace(' ', '_').replace('&', 'and')}.xlsx", r)
    res = ins.build(root)
    body = "Hello team,\n\nThe monthly reinsurance portfolio MI pack for 30 Sep 2026 is attached.\n\nKey points:\n" + \
           "\n".join(f"- {x}" for x in res["insights"]) + "\n\nRegional packs are in the shared folder.\n\nPortfolio Analytics"
    (root / "distribution" / "email_summary.txt").write_text(body)
    email(main_pack, body, root / "distribution" / "Monthly_MI_Email.eml")
    print("Reports written:", main_pack.name, "+ 5 regional packs + email draft")


if __name__ == "__main__":
    main()
